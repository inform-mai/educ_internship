import os
import pdfplumber
from docx import Document


def read_txt(file_path):
    with open(file_path, "r", encoding="utf-8") as f:
        text = f.read()
    return text


def table_to_markdown(table):
    # превращаем таблицу (список списков) в текст markdown
    if not table or len(table) == 0:
        return ""

    lines = []

    header = table[0]
    header = [cell if cell is not None else "" for cell in header]
    lines.append("| " + " | ".join(header) + " |")
    lines.append("| " + " | ".join(["---"] * len(header)) + " |")

    for row in table[1:]:
        row = [cell if cell is not None else "" for cell in row]
        lines.append("| " + " | ".join(row) + " |")

    return "\n".join(lines)


def read_pdf(file_path):
    full_text = ""

    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()

            if page_text is not None:
                full_text += page_text + "\n"

            tables = page.extract_tables()
            for table in tables:
                full_text += "\n" + table_to_markdown(table) + "\n"

    return full_text


def read_docx(file_path):
    # берём только текст, картинки не трогаем
    doc = Document(file_path)
    full_text = ""

    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            full_text += paragraph.text + "\n"

    # таблицы в word тоже переводим в markdown
    for table in doc.tables:
        rows = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            rows.append(cells)

        full_text += "\n" + table_to_markdown(rows) + "\n"

    return full_text


def parse_file(file_path):
    # кидаем файл, получаем текст
    # работает с .txt, .pdf и .docx

    if not os.path.exists(file_path):
        print(f"файл {file_path} не найден")
        return None

    if file_path.endswith(".txt"):
        text = read_txt(file_path)

    elif file_path.endswith(".pdf"):
        text = read_pdf(file_path)

    elif file_path.endswith(".docx"):
        text = read_docx(file_path)

    else:
        print("нужен .txt, .pdf или .docx")
        return None

    print(f"{file_path} прочитан, {len(text)} символов")
    return text